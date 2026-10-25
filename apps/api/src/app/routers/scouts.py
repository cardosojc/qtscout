import io
import math
from collections.abc import Sequence
from dataclasses import asdict
from datetime import UTC, datetime
from typing import Annotated, Any

from fastapi import APIRouter, Body, Depends, File, HTTPException, Query, UploadFile
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import delete, func, select
from sqlalchemy.dialects.sqlite import insert as sqlite_insert
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.camp_nights import CampNightRow, map_camp_night_row, milestones_reached
from app.core.document_utils import db_timestamp
from app.core.scout_utils import NIGHTS_BADGE_COUNTS, is_nights_badge_count
from app.core.siie_import import map_row
from app.core.text_sort import name_key
from app.db import atomic_batch, get_session
from app.deps import AdminUser, CurrentUser
from app.models import Profile, Scout, ScoutCampNight, ScoutNightsBadge
from app.models._columns import utcnow
from app.models.enums import OrdemSection
from app.models.ids import cuid_id
from app.schemas.scout import (
    CampNightMilestoneOut,
    CampNightOut,
    NightsBadgeOut,
    ScoutCampNightsOut,
    ScoutOut,
)

router = APIRouter(prefix="/scouts", tags=["scouts"])

# Rows per multi-row INSERT: 10 params each, under D1's 100-parameter limit.
_CAMP_ROWS_PER_INSERT = 9


async def _read_sheet(file: UploadFile | None) -> list[dict[str, Any]]:
    """First worksheet of an uploaded xlsx as dicts keyed by the header row."""
    if file is None:
        raise HTTPException(status_code=400, detail="Ficheiro não fornecido")
    try:
        content = await file.read()
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        if not workbook.worksheets:
            raise HTTPException(status_code=400, detail="Folha vazia")
        it = workbook.worksheets[0].iter_rows(values_only=True)
        header = list(next(it, []))
        return [
            dict(zip(header, row, strict=False)) for row in it if any(v is not None for v in row)
        ]
    except HTTPException:
        raise
    except Exception as err:  # noqa: BLE001
        raise HTTPException(
            status_code=400, detail=f"Não foi possível ler o ficheiro: {err}"
        ) from err


# camelCase request key -> snake_case model attribute (optional string fields).
_OPTIONAL_FIELDS = {
    "numeroAssociado": "numero_associado",
    "sexo": "sexo",
    "cc": "cc",
    "nif": "nif",
    "email": "email",
    "telefone": "telefone",
    "telemovel": "telemovel",
    "morada": "morada",
    "localidade": "localidade",
    "codigoPostal": "codigo_postal",
    "paiNome": "pai_nome",
    "paiTelefone": "pai_telefone",
    "paiEmail": "pai_email",
    "maeNome": "mae_nome",
    "maeTelefone": "mae_telefone",
    "maeEmail": "mae_email",
    "encarregadoNome": "encarregado_nome",
    "encarregadoTelefone": "encarregado_telefone",
    "encarregadoEmail": "encarregado_email",
}

_VALID_SECTIONS = {s.value for s in OrdemSection}


def _optional_string(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    return value.strip() or None


def _parse_date(value: Any) -> datetime | None:
    if not isinstance(value, str) or not value:
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed.astimezone(UTC).replace(tzinfo=None) if parsed.tzinfo else parsed


@router.get("")
async def list_scouts(
    user: CurrentUser,
    session: Annotated[Session, Depends(get_session)],
    section: Annotated[str | None, Query()] = None,
    includeInactive: Annotated[bool, Query()] = False,  # noqa: N803
) -> dict[str, list[ScoutOut]]:
    conditions = []
    if section and section in _VALID_SECTIONS:
        conditions.append(Scout.section == OrdemSection(section))
    if not includeInactive:
        conditions.append(Scout.active.is_(True))
    rows = (session.scalars(select(Scout).where(*conditions))).all()
    # Section (unassigned last, as before on Postgres), then surname, first name.
    rows = sorted(
        rows,
        key=lambda s: (
            s.section is None,
            s.section or "",
            name_key(s.last_name),
            name_key(s.first_name),
        ),
    )
    return {"scouts": [ScoutOut.model_validate(s) for s in rows]}


@router.post("")
async def create_scout(
    admin: AdminUser,
    session: Annotated[Session, Depends(get_session)],
    body: Annotated[dict[str, Any], Body()],
) -> dict[str, ScoutOut]:
    first_name = body["firstName"].strip() if isinstance(body.get("firstName"), str) else ""
    last_name = body["lastName"].strip() if isinstance(body.get("lastName"), str) else ""
    if not first_name or not last_name:
        raise HTTPException(status_code=400, detail="Nome e apelido obrigatórios")

    date_of_birth = _parse_date(body.get("dateOfBirth"))
    if not date_of_birth:
        raise HTTPException(status_code=400, detail="Data de nascimento inválida")

    section: OrdemSection | None = None
    raw_section = body.get("section")
    if raw_section is not None and raw_section != "":
        if raw_section not in _VALID_SECTIONS:
            raise HTTPException(status_code=400, detail="Secção inválida")
        section = OrdemSection(raw_section)

    joined_at = _parse_date(body.get("joinedAt")) or datetime.now(UTC).replace(tzinfo=None)
    raw_noites: Any = body.get("noitesCampoInicial")
    try:
        noites = float(raw_noites)
        noites_inicial = max(0, math.floor(noites)) if math.isfinite(noites) else 0
    except (TypeError, ValueError):
        noites_inicial = 0

    scout = Scout(
        first_name=first_name,
        last_name=last_name,
        date_of_birth=date_of_birth,
        section=section,
        joined_at=joined_at,
        noites_campo_inicial=noites_inicial,
    )
    for camel, attr in _OPTIONAL_FIELDS.items():
        setattr(scout, attr, _optional_string(body.get(camel)))

    session.add(scout)
    try:
        session.commit()
    except IntegrityError:
        session.rollback()
        raise HTTPException(status_code=409, detail="Nº de associado já existe") from None
    session.refresh(scout)
    return {"scout": ScoutOut.model_validate(scout)}


@router.post("/import")
async def import_scouts(
    admin: AdminUser,
    session: Annotated[Session, Depends(get_session)],
    file: Annotated[UploadFile | None, File()] = None,
) -> dict[str, Any]:
    records = await _read_sheet(file)

    summary: dict[str, Any] = {
        "total": len(records),
        "created": 0,
        "updated": 0,
        "linkedToProfile": 0,
        "errors": [],
    }

    valid = []
    for i, record in enumerate(records):
        mapped = map_row(record, i + 2)
        if not mapped.ok:
            summary["errors"].append(
                {"row": mapped.row, "nome": mapped.nome, "error": mapped.error}
            )
        else:
            valid.append(mapped)

    if not valid:
        return {"summary": summary}

    nins = [m.nin for m in valid]
    existing_nins = set(
        (
            session.scalars(select(Scout.numero_associado).where(Scout.numero_associado.in_(nins)))
        ).all()
    )
    profiles = (session.execute(select(Profile.id, Profile.email))).all()
    email_to_id = {email.lower(): pid for pid, email in profiles}

    for mapped in valid:
        assert mapped.payload is not None and mapped.nin is not None
        payload = mapped.payload
        was_existing = mapped.nin in existing_nins
        linked_id = email_to_id.get(payload.email.lower()) if payload.email else None
        fields = asdict(payload)
        if payload.joined_at is None:
            fields.pop("joined_at")
        try:
            if was_existing:
                scout = session.scalar(select(Scout).where(Scout.numero_associado == mapped.nin))
                if scout is not None:
                    for attr, value in fields.items():
                        setattr(scout, attr, value)
            else:
                if linked_id:
                    fields["profile_id"] = linked_id
                session.add(Scout(**fields))
            session.commit()
            if was_existing:
                summary["updated"] += 1
            else:
                summary["created"] += 1
                if linked_id:
                    summary["linkedToProfile"] += 1
        except IntegrityError:
            session.rollback()
            summary["errors"].append(
                {
                    "row": mapped.row,
                    "nome": f"{payload.first_name} {payload.last_name}".strip(),
                    "error": "Conflito de unicidade (já existe outro membro)",
                }
            )

    return {"summary": summary}


@router.post("/import-camp-nights")
async def import_camp_nights(
    admin: AdminUser,
    session: Annotated[Session, Depends(get_session)],
    file: Annotated[UploadFile | None, File()] = None,
) -> dict[str, Any]:
    """Import SIIE "Atividades por Elemento": each matched scout's rows are
    replaced (SIIE is the source of truth); unknown NINs (e.g. former members)
    are counted, not stored."""
    records = await _read_sheet(file)
    errors: list[dict[str, Any]] = []
    by_nin: dict[str, dict[str, CampNightRow]] = {}
    for i, record in enumerate(records):
        mapped = map_camp_night_row(record, i + 2)
        if not mapped.ok or mapped.payload is None:
            errors.append({"row": mapped.row, "error": mapped.error})
            continue
        # Last row wins if SIIE lists the same participation twice.
        by_nin.setdefault(mapped.payload.nin, {})[mapped.payload.activity_id] = mapped.payload

    scout_ids: dict[str, str] = {
        nin: scout_id
        for nin, scout_id in session.execute(
            select(Scout.numero_associado, Scout.id).where(Scout.numero_associado.is_not(None))
        ).all()
        if nin
    }
    now = db_timestamp(utcnow())
    matched = 0
    imported = 0
    for nin, rows in by_nin.items():
        scout_id = scout_ids.get(nin)
        if scout_id is None:
            continue
        matched += 1
        statements: list[tuple[str, Sequence[Any]]] = [
            ('DELETE FROM scout_camp_nights WHERE "scoutId" = ?', (scout_id,))
        ]
        values = list(rows.values())
        for start in range(0, len(values), _CAMP_ROWS_PER_INSERT):
            chunk = values[start : start + _CAMP_ROWS_PER_INSERT]
            placeholders = ", ".join("(?, ?, ?, ?, ?, ?, ?, ?, ?, ?)" for _ in chunk)
            params: list[Any] = []
            for r in chunk:
                params += [
                    cuid_id(),
                    scout_id,
                    r.activity_id,
                    r.description,
                    r.activity_type,
                    db_timestamp(r.start_date),
                    db_timestamp(r.end_date) if r.end_date else None,
                    r.nights,
                    r.camp_nights,
                    now,
                ]
            statements.append(
                (
                    'INSERT INTO scout_camp_nights (id, "scoutId", "activityId", description, '
                    '"activityType", "startDate", "endDate", nights, "campNights", "createdAt") '
                    f"VALUES {placeholders}",
                    params,
                )
            )
        atomic_batch(session, statements)
        imported += len(values)

    return {
        "summary": {
            "total": len(records),
            "matchedScouts": matched,
            "unmatchedMembers": len(by_nin) - matched,
            "imported": imported,
            "errors": errors[:50],
            "errorCount": len(errors),
        }
    }


@router.get("/camp-nights")
async def camp_night_totals(
    user: CurrentUser, session: Annotated[Session, Depends(get_session)]
) -> dict[str, dict[str, int]]:
    """Camp-night total per scout, for scouts with imported SIIE data."""
    rows = session.execute(
        select(ScoutCampNight.scout_id, func.sum(ScoutCampNight.camp_nights)).group_by(
            ScoutCampNight.scout_id
        )
    ).all()
    return {"totals": {scout_id: int(total or 0) for scout_id, total in rows}}


@router.get("/{scout_id}/camp-nights")
async def scout_camp_nights(
    scout_id: str, user: CurrentUser, session: Annotated[Session, Depends(get_session)]
) -> ScoutCampNightsOut:
    rows = session.scalars(
        select(ScoutCampNight)
        .where(ScoutCampNight.scout_id == scout_id)
        .order_by(ScoutCampNight.start_date.desc())
    ).all()
    recorded = {
        b.count: b.awarded_at
        for b in session.scalars(
            select(ScoutNightsBadge).where(ScoutNightsBadge.scout_id == scout_id)
        ).all()
    }
    milestones = [
        CampNightMilestoneOut(
            count=m.count,
            reached_on=m.reached_on,
            activity_id=m.activity_id,
            description=m.description,
            recorded_on=recorded.get(m.count),
        )
        for m in milestones_reached(list(rows))
    ]
    return ScoutCampNightsOut(
        total=sum(r.camp_nights for r in rows),
        activities=[CampNightOut.model_validate(r) for r in rows],
        milestones=milestones,
        imported_at=max((r.created_at for r in rows), default=None),
    )


@router.get("/{scout_id}")
async def get_scout(
    scout_id: str, user: CurrentUser, session: Annotated[Session, Depends(get_session)]
) -> dict[str, ScoutOut]:
    scout = session.scalar(select(Scout).where(Scout.id == scout_id))
    if scout is None:
        raise HTTPException(status_code=404, detail="Membro não encontrado")
    return {"scout": ScoutOut.model_validate(scout)}


@router.patch("/{scout_id}")
async def update_scout(
    scout_id: str,
    admin: AdminUser,
    session: Annotated[Session, Depends(get_session)],
    body: Annotated[dict[str, Any], Body()],
) -> dict[str, ScoutOut]:
    scout = session.scalar(select(Scout).where(Scout.id == scout_id))
    if scout is None:
        raise HTTPException(status_code=404, detail="Membro não encontrado")

    if "firstName" in body and isinstance(body["firstName"], str):
        if not body["firstName"].strip():
            raise HTTPException(status_code=400, detail="Nome obrigatório")
        scout.first_name = body["firstName"].strip()
    if "lastName" in body and isinstance(body["lastName"], str):
        if not body["lastName"].strip():
            raise HTTPException(status_code=400, detail="Apelido obrigatório")
        scout.last_name = body["lastName"].strip()
    if "dateOfBirth" in body:
        parsed = _parse_date(body["dateOfBirth"])
        if not parsed:
            raise HTTPException(status_code=400, detail="Data de nascimento inválida")
        scout.date_of_birth = parsed
    if "section" in body:
        raw = body["section"]
        if raw is None or raw == "":
            scout.section = None
        elif raw in _VALID_SECTIONS:
            scout.section = OrdemSection(raw)
        else:
            raise HTTPException(status_code=400, detail="Secção inválida")
    if "joinedAt" in body:
        parsed = _parse_date(body["joinedAt"])
        if not parsed:
            raise HTTPException(status_code=400, detail="Data de admissão inválida")
        scout.joined_at = parsed
    if "active" in body:
        scout.active = bool(body["active"])
    if "noitesCampoInicial" in body:
        try:
            n = float(body["noitesCampoInicial"])
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Noites de campo inválido") from None
        if not math.isfinite(n) or n < 0:
            raise HTTPException(status_code=400, detail="Noites de campo inválido")
        scout.noites_campo_inicial = math.floor(n)
    for camel, attr in _OPTIONAL_FIELDS.items():
        if camel in body:
            setattr(scout, attr, _optional_string(body[camel]))

    try:
        session.commit()
    except IntegrityError:
        session.rollback()
        raise HTTPException(status_code=409, detail="Nº de associado já existe") from None
    session.refresh(scout)
    return {"scout": ScoutOut.model_validate(scout)}


@router.delete("/{scout_id}")
async def delete_scout(
    scout_id: str, admin: AdminUser, session: Annotated[Session, Depends(get_session)]
) -> dict[str, bool]:
    scout = session.scalar(select(Scout).where(Scout.id == scout_id))
    if scout is None:
        raise HTTPException(status_code=404, detail="Membro não encontrado")
    session.delete(scout)
    session.commit()
    return {"ok": True}


@router.get("/{scout_id}/nights-badges")
async def get_nights_badges(
    scout_id: str, user: CurrentUser, session: Annotated[Session, Depends(get_session)]
) -> dict[str, list[NightsBadgeOut]]:
    rows = (
        session.scalars(
            select(ScoutNightsBadge)
            .where(ScoutNightsBadge.scout_id == scout_id)
            .order_by(ScoutNightsBadge.count.asc())
        )
    ).all()
    return {"badges": [NightsBadgeOut.model_validate(b) for b in rows]}


@router.put("/{scout_id}/nights-badges")
async def put_nights_badges(
    scout_id: str,
    admin: AdminUser,
    session: Annotated[Session, Depends(get_session)],
    body: Annotated[dict[str, Any], Body()],
) -> dict[str, Any]:
    scout = session.scalar(select(Scout.id).where(Scout.id == scout_id))
    if scout is None:
        raise HTTPException(status_code=404, detail="Membro não encontrado")

    badges = body.get("badges")
    if not isinstance(badges, list):
        raise HTTPException(status_code=400, detail="Formato inválido")

    to_delete: list[int] = []
    to_upsert: list[tuple[int, datetime]] = []
    for raw in badges:
        if raw is None or not isinstance(raw, dict):
            continue
        count: Any = raw.get("count")
        if not is_nights_badge_count(count):
            raise HTTPException(status_code=400, detail=f"count inválido ({count})")
        awarded = raw.get("awardedAt")
        if awarded is None or awarded == "":
            to_delete.append(int(count))
            continue
        parsed = _parse_date(awarded)
        if not parsed:
            raise HTTPException(status_code=400, detail=f"Data inválida para {count}")
        to_upsert.append((int(count), parsed))

    if to_delete:
        session.execute(
            delete(ScoutNightsBadge).where(
                ScoutNightsBadge.scout_id == scout_id,
                ScoutNightsBadge.count.in_(to_delete),
            )
        )
    for count, awarded_at in to_upsert:
        stmt = (
            sqlite_insert(ScoutNightsBadge)
            .values(scout_id=scout_id, count=count, awarded_at=awarded_at)
            .on_conflict_do_update(
                index_elements=[ScoutNightsBadge.scout_id, ScoutNightsBadge.count],
                # Column objects, not attribute names: the column is "awardedAt".
                set_={
                    ScoutNightsBadge.awarded_at: awarded_at,
                    ScoutNightsBadge.updated_at: utcnow(),
                },
            )
        )
        session.execute(stmt)
    session.commit()

    rows = (
        session.scalars(
            select(ScoutNightsBadge)
            .where(ScoutNightsBadge.scout_id == scout_id)
            .order_by(ScoutNightsBadge.count.asc())
        )
    ).all()
    return {
        "badges": [NightsBadgeOut.model_validate(b) for b in rows],
        "allCounts": list(NIGHTS_BADGE_COUNTS),
    }
